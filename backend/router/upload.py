import io
import os
import uuid
from datetime import datetime

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.agent.parser import ExcelParser, ExcelParseError
from backend.config import settings
from backend.database import get_db
from backend.models import ClassModel, StudentModel
from backend.schemas import UploadResponse

router = APIRouter(prefix="/api/v1", tags=["upload"])


@router.get("/template")
async def download_template():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "评语模板"

    headers = [
        ("姓名", True),
        ("成绩", False),
        ("课堂表现", False),
        ("作业情况", False),
        ("关键词", False),
    ]
    header_fill = PatternFill(start_color="409EFF", fgColor="409EFF", fill_type="solid")
    optional_fill = PatternFill(start_color="E6A23C", fgColor="E6A23C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    example_font = Font(color="999999", italic=True)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for col_idx, (label, required) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill if required else optional_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

        note = "必填" if required else "选填"
        cell.comment = openpyxl.comments.Comment(note, "系统")

    example = ["例：张三", "95 / 优秀", "例：积极举手发言", "例：按时完成", "例：思维活跃"]
    for col_idx, val in enumerate(example, 1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.font = example_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    instruction_fill = PatternFill(start_color="F0F9EB", fgColor="F0F9EB", fill_type="solid")
    instruction_cell = ws.cell(row=4, column=1, value="说明：")
    instruction_cell.font = Font(bold=True, size=10)
    instruction_cell.fill = instruction_fill

    instructions = [
        "• 蓝色标题 = 必填列，橙色标题 = 选填列",
        "• 「成绩」支持分数(95)或等级(优秀/良好/及格/A/B)",
        "• 「关键词」可简短描述学生特点，如：思维活跃、不够自信等",
        "• 列名可使用同义词，系统会自动识别（如「分数」=「成绩」）",
    ]
    for i, text in enumerate(instructions):
        cell = ws.cell(row=5 + i, column=1, value=text)
        cell.font = Font(color="666666", size=10)
        cell.fill = instruction_fill

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template.xlsx; filename*=utf-8''%E8%AF%84%E8%AF%AD%E6%A8%A1%E6%9D%BF.xlsx"},
    )


@router.post("/upload-excel", response_model=UploadResponse)
async def upload_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls 格式")

    content = await file.read()
    if len(content) > settings.max_upload_size:
        raise HTTPException(status_code=400, detail="文件大小超过 10MB 限制")

    filename = f"{uuid.uuid4().hex}_{file.filename}"
    filepath = os.path.join(settings.upload_dir, filename)
    with open(filepath, "wb") as f:
        f.write(content)

    try:
        preview_students, columns = ExcelParser.parse_preview(filepath)
        all_students, _ = ExcelParser.parse(filepath)
    except ExcelParseError as e:
        os.remove(filepath)
        raise HTTPException(status_code=400, detail=str(e))

    class_ = ClassModel(
        name=f"班级_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
        semester=datetime.now().strftime('%Y-%m'),
        style=settings.default_style,
        created_at=datetime.now(),
    )
    db.add(class_)
    await db.flush()

    for s in all_students:
        student = StudentModel(
            class_id=class_.id,
            name=s.name,
            score=s.score,
            performance=s.performance,
            homework=s.homework,
            keywords=s.keywords,
        )
        db.add(student)

    await db.commit()

    return UploadResponse(
        class_id=class_.id,
        class_name=class_.name,
        total=len(all_students),
        preview=preview_students,
        columns=columns,
    )
